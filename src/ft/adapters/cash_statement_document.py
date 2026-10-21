"""一次读取平台账单的原始出现和文件声明，不去重或建立账务关系。"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from io import BytesIO, StringIO
import csv
import re
from xml.parsers import expat
from zipfile import ZipFile

from ft.domain.cash_source_matching import SourceCoverage
from ft.domain.import_time import normalize_statement_timestamp


class StatementDocumentError(ValueError):
    def __init__(self, code):
        self.code = code
        super().__init__(code)


@dataclass(frozen=True)
class RawCashRow:
    position: int
    payload: dict[str, str] | None
    error_code: str | None = None
    document_position: str | None = None


@dataclass(frozen=True)
class CashStatementDocument:
    source_type: str
    table_key: str
    parser_version: str
    rows: tuple[RawCashRow, ...]
    source_account_identifier: str | None
    source_account_label: str | None
    coverage: SourceCoverage
    declared_count: int | None
    observed_count: int
    source_currency: str | None = None


_REQUIRED = {
    "alipay": {"交易时间", "收/支", "金额", "交易对方"},
    "wechat": {"交易时间", "收/支", "金额(元)", "交易对方"},
}
_MAX_INPUT_BYTES = 100 * 1024 * 1024
_MAX_PARSED_BYTES = 25 * 1024 * 1024
_MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
_MAX_ROWS = 100_000
_MAX_PHYSICAL_ROWS = 200_000
_MAX_COLUMNS = 128
_DATE = r"\d{4}[-/]\d{2}[-/]\d{2}[ T]\d{2}:\d{2}:\d{2}(?:Z|[+-]\d{2}:\d{2})?"


def _label_values(lines, *labels, inline=False):
    values = []
    for line in lines:
        for label in labels:
            prefix = r"(?:^|(?<=[\]】,]))\s*" if inline else r"^\s*"
            for match in re.finditer(prefix + re.escape(label) + r"\s*[:：]\s*(?:\[([^\]]*)\]|【([^】]*)】|([^\r\n,]*))", line):
                values.append(next(value for value in match.groups() if value is not None).strip())
    return values


def _unanimous(values):
    distinct = set(values)
    return distinct.pop() if len(distinct) == 1 else None


def _declarations(lines, source, observed, *, usable):
    text = "\n".join(lines)
    count_values = re.findall(r"(?m)^\s*共\s*(\d[\d,]*)\s*笔(?:记录|交易)?", text)
    if any(len(value) > 12 for value in count_values):
        raise StatementDocumentError("statement_rows_too_large")
    count = _unanimous(int(value.replace(",", "")) for value in count_values)
    account = _unanimous(_label_values(lines, *({"alipay": ("支付宝账户", "支付宝账号", "账号"), "wechat": ("微信号",)}[source]))) or None
    label = account or _unanimous(_label_values(lines, "微信昵称")) or None
    if account and re.search(r"[*＊•●]", account):
        account = None
    filter_labels = ("导出类型",) if source == "wechat" else ("导出交易类型", "交易类型")
    export_type = _unanimous("all" if value in {"全部", "全部交易", "所有", "全部类型", "所有交易"} else value or None
        for value in _label_values(lines, *filter_labels, inline=True))
    dates = []
    for label_pattern in ("起始时间", "终止时间"):
        normalized = []
        for value in _label_values(lines, label_pattern, inline=True):
            try:
                normalized.append(normalize_statement_timestamp(value.replace("/", "-"), source=source) if re.fullmatch(_DATE, value) else None)
            except ValueError:
                normalized.append(None)
        dates.append(_unanimous(normalized))
    complete = (
        usable and count is not None and count == observed and export_type is not None
        and all(dates) and datetime.fromisoformat(dates[0]) < datetime.fromisoformat(dates[1])
    )
    return account, label, count, SourceCoverage(
        start=dates[0], end=dates[1], start_inclusive=True, end_inclusive=True,
        complete=bool(complete), evidence="platform-declared-range-filter-count-v1" if complete else "",
        filters={"export_type": export_type} if export_type is not None else {},
        strategy_version="platform-coverage-v1",
    )


def _collect(records, source, table_key, *, pad_trailing=False):
    header, metadata, rows = None, [], []
    parsed_bytes = 0
    for position, values in enumerate(records, 1):
        if position > _MAX_PHYSICAL_ROWS:
            raise StatementDocumentError("statement_rows_too_large")
        cells = [str(value) if value is not None else "" for value in values]
        if cells:
            cells[0] = cells[0].removeprefix("\ufeff")
        if header is None:
            if position > 50 or sum(len(line) for line in metadata) > 65_536:
                raise StatementDocumentError("statement_header_not_found")
            matching_sources = [key for key, required in _REQUIRED.items() if required.issubset(cells)]
            if source is None and len(matching_sources) > 1:
                raise StatementDocumentError("statement_header_invalid")
            detected_source = source or (matching_sources[0] if matching_sources else None)
            if detected_source in matching_sources:
                if len(cells) > _MAX_COLUMNS or len(set(cells)) != len(cells):
                    raise StatementDocumentError("statement_header_invalid")
                header = cells
                source = detected_source
            else:
                last_value = max((index + 1 for index, value in enumerate(cells) if value), default=0)
                metadata.append(",".join(cells[:last_value]))
            continue
        if not any(cell.strip() for cell in cells):
            continue
        if pad_trailing and len(cells) < len(header):
            cells.extend([""] * (len(header) - len(cells)))
        if len(rows) >= _MAX_ROWS:
            raise StatementDocumentError("statement_rows_too_large")
        parsed_bytes += sum(len(cell.encode("utf-8")) for cell in cells)
        if parsed_bytes > _MAX_PARSED_BYTES:
            raise StatementDocumentError("statement_rows_too_large")
        if len(cells) != len(header):
            rows.append(RawCashRow(position, None, "row_shape_invalid"))
        else:
            rows.append(RawCashRow(position, dict(zip(header, cells))))
    if header is None:
        raise StatementDocumentError("statement_header_not_found")
    account, label, count, coverage = _declarations(metadata, source, len(rows), usable=all(row.error_code is None for row in rows))
    return CashStatementDocument(
        source_type=source, table_key=table_key, parser_version="platform-rows-v1",
        rows=tuple(rows), source_account_identifier=account, source_account_label=label,
        coverage=coverage, declared_count=count, observed_count=len(rows),
    )


def _preflight_workbook(archive):
    """在工作簿分配字符串／展开稀疏行前验证解压预算、XML 与实际坐标。"""
    members = archive.infolist()
    if len(members) > 1000 or len({member.filename for member in members}) != len(members) or sum(member.file_size for member in members) > _MAX_UNCOMPRESSED_BYTES:
        raise StatementDocumentError("statement_rows_too_large")
    for member in members:
        if member.is_dir():
            if member.file_size:
                raise StatementDocumentError("statement_unreadable")
            continue
        if not member.filename.endswith((".xml", ".rels")):
            # 平台表格只支持 XML 部件；关系目标不能借后缀或目录绕过预检。
            raise StatementDocumentError("statement_unreadable")
        worksheet = False
        depth, row_count, cell_count, current_row, last_column = 0, 0, 0, 0, 0

        def forbidden(*_args):
            raise StatementDocumentError("statement_unreadable")

        def start(name, attributes):
            nonlocal depth, row_count, cell_count, current_row, last_column, worksheet
            depth += 1
            if depth > 128:
                raise StatementDocumentError("statement_unreadable")
            local_name = name.rsplit("}", 1)[-1]
            if (depth == 1 and local_name == "worksheet") or name == "http://schemas.openxmlformats.org/spreadsheetml/2006/main}row":
                worksheet = True
            if not worksheet:
                return
            if local_name == "row":
                row_count += 1
                cell_count, last_column = 0, 0
                index = attributes.get("r", str(current_row + 1))
                if not index.isascii() or not index.isdigit() or len(index) > 7 or not 1 <= int(index) <= _MAX_PHYSICAL_ROWS or row_count > _MAX_PHYSICAL_ROWS:
                    raise StatementDocumentError("statement_rows_too_large")
                if int(index) <= current_row:
                    raise StatementDocumentError("statement_unreadable")
                current_row = int(index)
            elif local_name == "c":
                cell_count += 1
                coordinate = attributes.get("r")
                match = re.fullmatch(r"([A-Z]{1,3})([0-9]{1,7})", coordinate) if coordinate else None
                if (coordinate and not match) or cell_count > _MAX_COLUMNS:
                    raise StatementDocumentError("statement_header_invalid")
                column = last_column + 1
                if match:
                    column = 0
                    for letter in match[1]:
                        column = column * 26 + ord(letter) - ord("A") + 1
                    if column > _MAX_COLUMNS:
                        raise StatementDocumentError("statement_header_invalid")
                    if not 1 <= int(match[2]) <= _MAX_PHYSICAL_ROWS:
                        raise StatementDocumentError("statement_rows_too_large")
                    if int(match[2]) != current_row:
                        raise StatementDocumentError("statement_unreadable")
                if current_row == 0 or column <= last_column:
                    raise StatementDocumentError("statement_unreadable")
                last_column = column

        def end(_name):
            nonlocal depth
            depth -= 1

        parser = expat.ParserCreate(namespace_separator="}")
        parser.StartDoctypeDeclHandler = forbidden
        parser.EntityDeclHandler = forbidden
        parser.ExternalEntityRefHandler = forbidden
        parser.StartElementHandler, parser.EndElementHandler = start, end
        with archive.open(member) as stream:
            while chunk := stream.read(65536):
                parser.Parse(chunk, False)
            parser.Parse(b"", True)


def read_platform_document(content: bytes, *, source: str | None = None) -> CashStatementDocument:
    if source is not None and source not in _REQUIRED:
        raise StatementDocumentError("statement_source_unsupported")
    if not isinstance(content, bytes) or not content or len(content) > _MAX_INPUT_BYTES:
        raise StatementDocumentError("statement_input_invalid")
    if content.startswith(b"PK"):
        import openpyxl
        try:
            with BytesIO(content) as stream:
                with ZipFile(stream) as archive:
                    _preflight_workbook(archive)
                stream.seek(0)
                workbook = openpyxl.load_workbook(stream, read_only=True, data_only=False, keep_links=False)
                try:
                    sheet = workbook.active
                    if sheet is None:
                        raise StatementDocumentError("statement_header_invalid")
                    sheet.reset_dimensions()
                    return _collect(sheet.iter_rows(values_only=True), source, f"sheet:{workbook.index(sheet)}", pad_trailing=True)
                finally:
                    workbook.close()
        except StatementDocumentError:
            raise
        except Exception:
            raise StatementDocumentError("statement_unreadable") from None
    text = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeError:
            continue
    if text is None:
        raise StatementDocumentError("statement_encoding_unsupported")
    try:
        with StringIO(text, newline="") as stream:
            return _collect(csv.reader(stream), source, "csv:0")
    except csv.Error:
        raise StatementDocumentError("statement_unreadable") from None
