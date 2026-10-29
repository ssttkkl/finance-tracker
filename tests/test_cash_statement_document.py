import csv
from io import BytesIO, StringIO

import pytest


ALIPAY_HEADER = ["交易时间", "交易分类", "交易对方", "对方账号", "商品说明", "收/支", "金额", "收/付款方式", "交易状态", "交易订单号", "商家订单号"]
WECHAT_HEADER = ["交易时间", "交易类型", "交易对方", "商品", "收/支", "金额(元)", "支付方式", "当前状态", "交易单号", "商户单号", "备注"]


def _csv(rows, *, count=None, export_type="全部", account="fixture-owner@example.test"):
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(["支付宝交易记录明细查询"])
    writer.writerow([f"支付宝账户：[{account}]"])
    writer.writerow(["起始时间：[2026-08-01 00:00:00] 终止时间：[2026-08-31 23:59:59]"])
    if export_type is not None:
        writer.writerow([f"交易类型：[{export_type}]"])
    if count is not None:
        writer.writerow([f"共{count}笔记录"])
    writer.writerow(ALIPAY_HEADER)
    writer.writerows(rows)
    return output.getvalue().encode("utf-8-sig")


def _alipay_row(**changes):
    fields = dict(zip(ALIPAY_HEADER, ["2026-08-15 10:00:00", "日用百货", "测试商户", "account", "  原始描述\n第二行  ", "支出", "20.00", "测试银行卡", "交易成功", "0000123", "merchant-1"]))
    fields.update(changes)
    return [fields[key] for key in ALIPAY_HEADER]


def _xlsx(rows, *, count=1, header=WECHAT_HEADER):
    import openpyxl
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in [["微信支付账单明细"], ["微信昵称：[测试昵称]"],
                ["起始时间：[2026-08-01 00:00:00]"], ["终止时间：[2026-08-31 23:59:59]"],
                ["导出类型：[全部]"], [f"共{count}笔记录"],
                ["说明：请核对申请导出账单的微信号，不代表本人微信号字段"], header, *rows]:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _read(content, source="alipay"):
    from ft.adapters.cash_statement_document import read_platform_document
    return read_platform_document(content, source=source)


def test_csv_preserves_every_occurrence_and_exact_native_fields():
    row = _alipay_row()
    document = _read(_csv([row, row], count=2))
    assert len(document.rows) == 2
    assert document.rows[0].position != document.rows[1].position
    assert document.rows[0].payload == document.rows[1].payload
    assert document.rows[0].payload["商品说明"] == "  原始描述\n第二行  "
    assert document.rows[0].payload["交易订单号"] == "0000123"
    assert document.declared_count == document.observed_count == 2
    assert document.coverage.complete
    assert document.source_account_identifier == "fixture-owner@example.test"
    assert document.coverage.contains("2026-08-31T15:59:59Z")


@pytest.mark.parametrize("count,export_type", [(None, "全部"), (2, "全部"), (1, None)])
def test_missing_declaration_or_wrong_count_never_infers_complete_from_transactions(count, export_type):
    document = _read(_csv([_alipay_row()], count=count, export_type=export_type))
    assert not document.coverage.complete


def test_different_declared_filters_are_not_comparable():
    all_rows = _read(_csv([_alipay_row()], count=1))
    filtered = _read(_csv([_alipay_row()], count=1, export_type="支出"))
    assert filtered.coverage.complete
    assert not filtered.coverage.comparable(all_rows.coverage, "2026-08-15T00:00:00Z")


def test_alipay_combined_metadata_line_retains_its_declared_filter():
    content = _csv([_alipay_row()], count=1, export_type=None)
    old = "起始时间：[2026-08-01 00:00:00] 终止时间：[2026-08-31 23:59:59]"
    content = content.replace(old.encode(), (old + " 交易类型：[全部]").encode())
    assert _read(content).coverage.complete


def test_alipay_export_filter_label_is_recognized_without_matching_arbitrary_prose():
    content = _csv([_alipay_row()], count=1).replace("交易类型：[全部]".encode(), "导出交易类型：[全部]".encode())
    assert _read(content).coverage.complete


def test_invalid_row_shape_isolated_and_other_rows_preserved():
    good = _alipay_row()
    document = _read(_csv([good, good[:-1], good], count=3))
    assert len(document.rows) == 3
    assert [row.error_code for row in document.rows] == [None, "row_shape_invalid", None]
    assert document.rows[1].payload is None
    assert document.observed_count == 3
    assert document.declared_count == 3
    assert not document.coverage.complete
    assert not document.coverage.evidence


def test_separator_like_bad_row_cannot_truncate_later_transactions():
    document = _read(_csv([_alipay_row(), ["-----------"], _alipay_row()], count=3))
    assert len(document.rows) == 3
    assert document.rows[-1].payload["交易订单号"] == "0000123"


def test_amount_errors_are_not_converted_to_zero_or_silently_skipped_by_raw_reader():
    document = _read(_csv([_alipay_row(金额="invalid"), _alipay_row(金额="0")], count=2))
    assert document.rows[0].payload["金额"] == "invalid"
    assert document.rows[1].payload["金额"] == "0"


def test_wechat_nickname_and_explanatory_account_word_are_not_owner_identity():
    row = ["2026-08-15 10:00:00", "商户消费", "测试店", "商品原文", "支出", "20.00", "零钱", "支付成功", "000123", "merchant", "备注原文"]
    document = _read(_xlsx([row]), "wechat")
    assert document.source_account_identifier is None
    assert document.source_account_label == "测试昵称"
    assert document.source_nickname == "测试昵称"
    assert document.rows[0].payload["备注"] == "备注原文"
    assert document.coverage.complete


def test_masked_account_is_only_a_label_not_a_source_identity():
    document = _read(_csv([_alipay_row()], count=1, account="138****1234"))
    assert document.source_account_identifier is None
    assert document.source_account_label == "138****1234"


def test_duplicate_header_fails_safely_without_source_values():
    from ft.adapters.cash_statement_document import StatementDocumentError
    content = _csv([_alipay_row()], count=1).replace("商品说明".encode(), "交易对方".encode())
    with pytest.raises(StatementDocumentError, match="statement_header_invalid") as error:
        _read(content)
    assert "fixture-owner" not in str(error.value)


def test_document_never_retains_upload_bytes_or_open_workbook():
    from dataclasses import asdict
    row = ["2026-08-15 10:00:00", "商户消费", "店", "商品", "支出", "20", "零钱", "支付成功", "id", "merchant", "备注"]
    document = _read(_xlsx([row]), "wechat")
    def inspect(value):
        assert not isinstance(value, (bytes, bytearray, BytesIO))
        if isinstance(value, dict):
            for item in value.values(): inspect(item)
        elif isinstance(value, (list, tuple)):
            for item in value: inspect(item)
    inspect(asdict(document))


def test_summary_subtotal_is_not_used_as_the_statement_record_count():
    content = _csv([_alipay_row()], count=None)
    content = content.replace("交易类型：[全部]".encode(), "交易类型：[全部]\r\n收入共1笔记录".encode())
    assert _read(content).declared_count is None
    assert not _read(content).coverage.complete


def test_workbook_expansion_is_bounded_before_openpyxl_reads_it(monkeypatch):
    import zipfile
    import openpyxl
    from ft.adapters import cash_statement_document as adapter
    stream = BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/sharedStrings.xml", "a" * 2048)
    monkeypatch.setattr(adapter, "_MAX_UNCOMPRESSED_BYTES", 1024)
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *_a, **_kw: pytest.fail("oversized archive must be rejected before workbook load"))
    with pytest.raises(adapter.StatementDocumentError, match="statement_rows_too_large"):
        _read(stream.getvalue(), "wechat")


def _rewrite_xlsx(content, change):
    from zipfile import ZipFile, ZIP_DEFLATED
    output = BytesIO()
    with ZipFile(BytesIO(content)) as original, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for member in original.infolist():
            target.writestr(member.filename, change(member.filename, original.read(member)))
    return output.getvalue()


def _wechat_row():
    return ["2026-08-15 10:00:00", "商户消费", "测试店", "商品", "支出", "20", "零钱", "支付成功", "000123", "merchant", "备注"]


def test_underreported_sheet_dimension_cannot_hide_later_rows():
    import re
    content = _rewrite_xlsx(_xlsx([_wechat_row(), _wechat_row()], count=2),
        lambda name, body: re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', body) if name == "xl/worksheets/sheet1.xml" else body)
    document = _read(content, "wechat")
    assert len(document.rows) == 2
    assert document.coverage.complete


def test_sparse_sheet_row_index_rejected_before_workbook_loading(monkeypatch):
    import openpyxl
    from ft.adapters import cash_statement_document as adapter
    content = _rewrite_xlsx(_xlsx([_wechat_row()]),
        lambda name, body: body.replace(b"</sheetData>", b'<row r="2000"></row></sheetData>') if name == "xl/worksheets/sheet1.xml" else body)
    monkeypatch.setattr(adapter, "_MAX_PHYSICAL_ROWS", 1000, raising=False)
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *_a, **_kw: pytest.fail("sparse indices must be checked before iteration"))
    with pytest.raises(adapter.StatementDocumentError, match="statement_rows_too_large"):
        _read(content, "wechat")


def test_physical_csv_rows_are_bounded_even_when_empty(monkeypatch):
    from ft.adapters import cash_statement_document as adapter
    content = _csv([_alipay_row()], count=1) + b"\n" * 100
    monkeypatch.setattr(adapter, "_MAX_PHYSICAL_ROWS", 20, raising=False)
    with pytest.raises(adapter.StatementDocumentError, match="statement_rows_too_large"):
        _read(content)


def test_out_of_budget_cell_cannot_be_hidden_by_sheet_dimensions():
    from ft.adapters.cash_statement_document import StatementDocumentError
    content = _rewrite_xlsx(_xlsx([_wechat_row()]),
        lambda name, body: body.replace(b"</sheetData>", b'<row r="10"><c r="ZZZZZZ10" t="n"><v>1</v></c></row></sheetData>') if name == "xl/worksheets/sheet1.xml" else body)
    with pytest.raises(StatementDocumentError, match="statement_header_invalid"):
        _read(content, "wechat")


def test_workbook_xml_dtd_rejected_before_expansion(monkeypatch):
    import openpyxl
    from ft.adapters.cash_statement_document import StatementDocumentError
    content = _rewrite_xlsx(_xlsx([_wechat_row()]),
        lambda name, body: b'<!DOCTYPE worksheet [<!ENTITY x "unsafe">]>' + body if name == "xl/worksheets/sheet1.xml" else body)
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *_a, **_kw: pytest.fail("DTD must be rejected before workbook load"))
    with pytest.raises(StatementDocumentError, match="statement_unreadable"):
        _read(content, "wechat")


@pytest.mark.parametrize("extra,attribute", [
    ("支付宝账户：[second@example.test]", "source_account_identifier"),
    ("共999笔记录", "declared_count"),
])
def test_conflicting_declarations_never_choose_first(extra, attribute):
    content = _csv([_alipay_row()], count=1).replace("支付宝交易记录明细查询".encode(), ("支付宝交易记录明细查询\r\n" + extra).encode())
    document = _read(content)
    assert getattr(document, attribute) is None
    if attribute == "declared_count":
        assert not document.coverage.complete


@pytest.mark.parametrize("extra", ["交易类型：[支出]", "起始时间：[2026-08-02 00:00:00]", "终止时间：[2026-09-01 00:00:00]"])
def test_conflicting_coverage_fields_are_unknown(extra):
    content = _csv([_alipay_row()], count=1).replace("支付宝交易记录明细查询".encode(), ("支付宝交易记录明细查询\r\n" + extra).encode())
    assert not _read(content).coverage.complete


def test_consistent_repeated_declarations_are_accepted():
    content = _csv([_alipay_row()], count=1).replace("支付宝交易记录明细查询".encode(), "支付宝交易记录明细查询\r\n共1笔记录\r\n交易类型：[全部交易]".encode())
    assert _read(content).coverage.complete


def test_adjacent_conflicting_fields_on_one_line_are_not_missed():
    content = _csv([_alipay_row()], count=1).replace("交易类型：[全部]".encode(), "交易类型：[全部]交易类型：[支出]".encode())
    assert not _read(content).coverage.complete


def test_xlsx_implicit_trailing_empty_cells_are_valid_fields():
    document = _read(_xlsx([_wechat_row()[:-1]]), "wechat")
    assert document.rows[0].error_code is None
    assert document.rows[0].payload["备注"] == ""
    assert document.coverage.complete


def test_wechat_transaction_id_column_is_required():
    from ft.adapters.cash_statement_document import StatementDocumentError
    header = WECHAT_HEADER[:8] + WECHAT_HEADER[9:]
    row = _wechat_row()[:8] + _wechat_row()[9:]
    with pytest.raises(StatementDocumentError, match="statement_header_not_found"):
        _read(_xlsx([row], header=header), "wechat")


@pytest.mark.parametrize("fragment", [
    b'<row r="9"><c r="A9" t="n"><v>1</v></c></row>',
    b'<row r="10"><c r="A9" t="n"><v>1</v></c></row>',
    b'<row r="10"><c r="B10" t="n"><v>1</v></c><c r="A10" t="n"><v>2</v></c></row>',
])
def test_corrupt_sheet_coordinates_cannot_silently_hide_occurrences(fragment):
    from ft.adapters.cash_statement_document import StatementDocumentError
    content = _rewrite_xlsx(_xlsx([_wechat_row()]),
        lambda name, body: body.replace(b"</sheetData>", fragment + b"</sheetData>") if name == "xl/worksheets/sheet1.xml" else body)
    with pytest.raises(StatementDocumentError, match="statement_unreadable"):
        _read(content, "wechat")


@pytest.mark.parametrize("target", ["xl/custom.bin", "xl/custom.xml", "xl/custom/", "xl/custom.xml/"])
def test_relationship_targets_cannot_bypass_xml_or_coordinate_checks(target):
    from zipfile import ZipFile, ZIP_DEFLATED
    from ft.adapters.cash_statement_document import StatementDocumentError
    output = BytesIO()
    with ZipFile(BytesIO(_xlsx([_wechat_row()]))) as original, ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for member in original.infolist():
            body = original.read(member).replace(b"/xl/worksheets/sheet1.xml", ("/" + target).encode())
            name = member.filename
            if name == "xl/worksheets/sheet1.xml":
                name = target
                if target.endswith((".bin", "/")):
                    body = b'<!DOCTYPE worksheet [<!ENTITY x "expanded">]>' + body
                else:
                    body = body.replace(b"</sheetData>", b'<row r="10"/><row r="9"/></sheetData>')
            archive.writestr(name, body)
    with pytest.raises(StatementDocumentError, match="statement_unreadable"):
        _read(output.getvalue(), "wechat")


@pytest.mark.parametrize("content,source", [(_csv([_alipay_row()], count=1), "alipay"), (_xlsx([_wechat_row()]), "wechat")], ids=["alipay-csv", "wechat-xlsx"])
def test_source_can_be_identified_from_the_same_header_read(content, source):
    document = _read(content, None)
    assert document.source_type == source
    assert document.observed_count == 1
    assert document.coverage.complete


def test_ambiguous_source_header_is_not_guessed():
    from ft.adapters.cash_statement_document import StatementDocumentError
    content = _csv([_alipay_row()], count=1).replace("商家订单号".encode(), "金额(元)".encode()).replace("交易订单号".encode(), "交易单号".encode())
    with pytest.raises(StatementDocumentError, match="statement_header_invalid"):
        _read(content, None)


def test_alipay_transaction_order_number_header_is_required():
    from ft.adapters.cash_statement_document import StatementDocumentError
    content = _csv([_alipay_row()], count=1).replace("交易订单号".encode(), "账单编号".encode())
    with pytest.raises(StatementDocumentError):
        _read(content, "alipay")
